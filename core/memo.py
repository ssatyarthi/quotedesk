"""Markdown memo -> .docx (headings, bullets, tables, paragraphs)."""
import io
import re


def md_to_docx(title: str, md: str) -> bytes:
    from docx import Document
    doc = Document()
    doc.add_heading(title, level=0)
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|\s*:?-", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|\s*:?-", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            t.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    t.cell(r, c).text = val.replace("**", "")
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            doc.add_heading(m.group(2), level=min(len(m.group(1)), 3))
        elif re.match(r"^\s*[-*]\s+", ln):
            doc.add_paragraph(re.sub(r"^\s*[-*]\s+", "", ln).replace("**", ""), style="List Bullet")
        elif re.match(r"^\s*\d+\.\s+", ln):
            doc.add_paragraph(re.sub(r"^\s*\d+\.\s+", "", ln).replace("**", ""), style="List Number")
        elif ln.strip():
            p = doc.add_paragraph()
            for part in re.split(r"(\*\*[^*]+\*\*)", ln):
                if part.startswith("**") and part.endswith("**"):
                    p.add_run(part[2:-2]).bold = True
                else:
                    p.add_run(part)
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
