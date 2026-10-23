"""
Turn whatever a vendor sent into content the model can read, while keeping
addresses (sheet!cell, page, paragraph) so every extracted number can be
traced back to where it came from.

No parsing of *meaning* happens here - only format conversion. All
interpretation is done by the model in extract.py.
"""
import base64
import io
import json
from pathlib import Path

from . import config

IMAGE_TYPES = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}


def load_manifest():
    return json.loads((config.RESPONSES / "manifest.json").read_text(encoding="utf-8"))


def vendor_files(vendor_id):
    d = config.RESPONSES / vendor_id
    return sorted(p for p in d.iterdir() if p.is_file())


def xlsx_to_text(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"--- sheet '{ws.title}' ---")
        merged = [str(r) for r in ws.merged_cells.ranges]
        if merged:
            out.append(f"(merged ranges: {', '.join(merged)})")
        for row in ws.iter_rows():
            cells = [f"{c.coordinate}={c.value!r}" for c in row if c.value not in (None, "")]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def docx_to_text(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    out = []
    for i, p in enumerate(doc.paragraphs, 1):
        if p.text.strip():
            style = p.style.name if p.style is not None else ""
            tag = f"[{style}] " if style.startswith("Heading") else ""
            out.append(f"¶{i}: {tag}{p.text}")
    for ti, t in enumerate(doc.tables, 1):
        for ri, r in enumerate(t.rows, 1):
            out.append(f"table{ti} row{ri}: " + " | ".join(c.text for c in r.cells))
    return "\n".join(out)


def image_block(path: Path):
    from PIL import Image
    im = Image.open(path)
    im = im.convert("RGB")
    im.thumbnail((1568, 1568))  # model's useful max resolution
    buf = io.BytesIO()
    im.save(buf, format="JPEG", quality=90)
    return {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg",
                                         "data": base64.b64encode(buf.getvalue()).decode()}}


def pdf_block(path: Path):
    return {"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                            "data": base64.b64encode(path.read_bytes()).decode()},
            "title": path.name}


def vendor_content_blocks(vendor_id):
    """Content blocks for one vendor's full response (email + every attachment)."""
    blocks = []
    for p in vendor_files(vendor_id):
        ext = p.suffix.lower()
        if ext == ".txt":
            blocks.append({"type": "text", "text": f"=== FILE: {p.name} (the vendor's email) ===\n{p.read_text(encoding='utf-8')}"})
        elif ext == ".xlsx":
            blocks.append({"type": "text", "text": f"=== FILE: {p.name} (Excel, every non-empty cell with its address) ===\n{xlsx_to_text(p)}"})
        elif ext == ".docx":
            blocks.append({"type": "text", "text": f"=== FILE: {p.name} (Word document, numbered paragraphs) ===\n{docx_to_text(p)}"})
        elif ext == ".pdf":
            blocks.append({"type": "text", "text": f"=== FILE: {p.name} (PDF follows; cite page numbers) ==="})
            blocks.append(pdf_block(p))
        elif ext in IMAGE_TYPES:
            blocks.append({"type": "text", "text": f"=== FILE: {p.name} (photo follows) ==="})
            blocks.append(image_block(p))
        else:
            blocks.append({"type": "text", "text": f"=== FILE: {p.name} - unsupported format, not read ==="})
    return blocks


def preview_text(path: Path) -> str:
    """Human preview used by the Inbox screen."""
    ext = path.suffix.lower()
    if ext == ".txt":
        return path.read_text(encoding="utf-8")
    if ext == ".xlsx":
        return xlsx_to_text(path)
    if ext == ".docx":
        return docx_to_text(path)
    if ext == ".pdf":
        try:
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                return "\n\n".join(f"--- page {i} ---\n{pg.extract_text() or ''}" for i, pg in enumerate(pdf.pages, 1))
        except Exception as e:
            return f"(could not preview PDF: {e})"
    return ""
